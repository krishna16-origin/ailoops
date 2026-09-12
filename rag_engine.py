"""
rag_engine.py
=================
Advanced Retrieval-Augmented Generation (RAG) over user-uploaded files and
images, shared by both Chat mode and Code mode.

Pipeline per uploaded file
---------------------------
  1. EXTRACT   Turn the raw upload into text (documents) or a rich visual
               analysis (images), using the right tool for that file type.
  2. CHUNK     Split the extracted text into overlapping, paragraph-aware
               chunks small enough to embed and retrieve precisely.
  3. EMBED     Turn each chunk into a vector, via a real NVIDIA embedding
               model when one is configured, with a deterministic local
               fallback so retrieval still works in demo / offline mode
               (mirrors the "DEMO fallback" pattern already used elsewhere
               in this app when NVIDIA_API_KEY isn't set).
  4. STORE     Keep chunks + vectors inside the existing per-session
               `sessions[session_id]` dict that app.py already threads
               through every request — no new global state, no external
               vector DB. Intentionally in-memory and session-scoped, same
               as conversation history.
  5. RETRIEVE  On every chat/code turn, embed the user's latest message and
               rank stored chunks with a hybrid score (cosine similarity +
               lexical overlap), returning the top-K most relevant chunks
               formatted as a context block for the model's system prompt.

Images get one extra step up front: a vision-capable model produces a
structured analysis (detailed description, any visible text/data/code,
notable entities) which is itself chunked/embedded/retrieved exactly like
any other document. That's what lets Chat and Code mode "see" and reason
over photos, screenshots, diagrams, and scanned pages a user attaches —
not just files with plain extractable text.
"""

import os
import io
import re
import math
import base64
import hashlib
import traceback
import asyncio
import uuid
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional, Tuple

import numpy as np

from langchain_core.messages import HumanMessage
from langchain_nvidia_ai_endpoints import ChatNVIDIA, NVIDIAEmbeddings

import session_store

try:
    import pypdf
except Exception:  # pragma: no cover - optional at import time
    pypdf = None

try:
    import docx as _docx  # python-docx
except Exception:  # pragma: no cover
    _docx = None

try:
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None


# ---------------------------------------------------------------------------
# Tunables
# ---------------------------------------------------------------------------
MAX_FILE_BYTES = 25 * 1024 * 1024          # 25MB per file
MAX_FILES_PER_SESSION = 60                 # generous cap on total uploads/session
CHUNK_CHARS = 1100                         # target characters per retrieval chunk
CHUNK_OVERLAP = 150                        # char overlap carried into the next chunk
TOP_K_DEFAULT = 6                          # chunks returned per retrieval call
MAX_CONTEXT_CHARS_PER_CHUNK = 900          # how much of a chunk is shown to the model
EMBED_DIM_FALLBACK = 512                   # dimensionality of the local hash embedding
VISION_IMAGE_MAX_DIM = 1600                # downscale before sending to the vision model
THUMB_MAX_DIM = 360                        # downscale for the small UI thumbnail

NVIDIA_EMBED_MODEL = os.getenv("NVIDIA_EMBED_MODEL", "nvidia/nv-embedqa-e5-v5")
VISION_MODEL = os.getenv("NVIDIA_VISION_MODEL", "meta/llama-3.2-90b-vision-instruct")
# Nemotron Parse is a separate, purpose-built OCR/document-parsing pass that
# runs on every image upload ALONGSIDE the general vision analysis above —
# it is never exposed as a selectable chat model. It exists purely so the
# text/tables/layout in a screenshot, scan, or photo of a document come
# through verbatim in the RAG context that gets handed to whichever chat
# model the user has actually picked (DeepSeek, Kimi, Gemma, a Nemotron chat
# model, Poolside, ...) — the extraction happens here, once, in the shared
# upload pipeline, so it "just works" no matter which model answers.
PARSE_MODEL = os.getenv("NVIDIA_PARSE_MODEL", "nvidia/nemotron-parse-2.0")
PARSE_IMAGE_MAX_DIM = 2200                 # OCR needs crisper text than general vision needs

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json",
             ".yaml", ".yml", ".xml", ".ini", ".cfg", ".toml", ".env"}
CODE_EXTS = {".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".cpp", ".cc", ".h",
             ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".sql",
             ".sh", ".bash", ".html", ".htm", ".css", ".scss", ".less", ".vue",
             ".svelte", ".dart", ".r", ".m", ".pl", ".lua"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".tif"}

ALL_SUPPORTED_EXTS = TEXT_EXTS | CODE_EXTS | PDF_EXTS | DOCX_EXTS | IMAGE_EXTS


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def classify(filename: str, content_type: str = "") -> str:
    """Return 'image' | 'pdf' | 'docx' | 'code' | 'text' | 'unsupported'."""
    ext = _ext(filename)
    if ext in IMAGE_EXTS or (content_type or "").startswith("image/"):
        return "image"
    if ext in PDF_EXTS:
        return "pdf"
    if ext in DOCX_EXTS:
        return "docx"
    if ext in CODE_EXTS:
        return "code"
    if ext in TEXT_EXTS:
        return "text"
    # Unknown extension but a text-ish content-type — still worth indexing.
    if (content_type or "").startswith("text/"):
        return "text"
    return "unsupported"


# ---------------------------------------------------------------------------
# Small local helpers (kept standalone so this module has no dependency on
# app.py — app.py imports this module, not the other way around).
# ---------------------------------------------------------------------------
def _coerce_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "".join(_coerce_text(v.get("text") or v.get("content") or "") if isinstance(v, dict) else _coerce_text(v) for v in value)
    if isinstance(value, dict):
        return _coerce_text(value.get("text") or value.get("content") or "")
    return str(value)


def _short_summary(text: str, limit: int = 220) -> str:
    clean = " ".join((text or "").split())
    if len(clean) <= limit:
        return clean
    return clean[: limit - 1].rstrip() + "…"


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
async def _extract_pdf(data: bytes) -> List[Tuple[Optional[int], str]]:
    """Returns a list of (page_number, page_text). Scanned/image-only pages
    come back with a short placeholder note rather than being silently
    dropped, so the model knows content may be missing rather than assuming
    the page was blank."""
    if pypdf is None:
        return [(None, "[PDF text extraction unavailable on this server — the 'pypdf' package is not installed.]")]

    def _sync():
        reader = pypdf.PdfReader(io.BytesIO(data))
        out: List[Tuple[Optional[int], str]] = []
        for i, page in enumerate(reader.pages):
            try:
                text = (page.extract_text() or "").strip()
            except Exception:
                text = ""
            if text:
                out.append((i + 1, text))
            else:
                out.append((i + 1, "[No extractable text on this page — likely a scanned/image page.]"))
        return out

    return await asyncio.to_thread(_sync)


async def _extract_docx(data: bytes) -> str:
    if _docx is None:
        return "[DOCX text extraction unavailable on this server — the 'python-docx' package is not installed.]"

    def _sync():
        d = _docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    return await asyncio.to_thread(_sync)


def _extract_plain(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="replace")


def _downscale_image(data: bytes, max_dim: int = VISION_IMAGE_MAX_DIM, quality: int = 85) -> bytes:
    if Image is None:
        return data
    try:
        img = Image.open(io.BytesIO(data))
        w, h = img.size
        # For large JPEGs, ask libjpeg to decode straight to (roughly) the target
        # size instead of decoding at full resolution and resizing afterwards.
        # This is a large, well-known speedup (often several times faster) for
        # multi-megapixel phone photos and is why a single upload could feel
        # slow before: we were fully decoding the original photo twice (once
        # for the small UI thumbnail, once again later for the vision model).
        try:
            img.draft("RGB", (max_dim, max_dim))
        except Exception:
            pass
        img = img.convert("RGB")
        w, h = img.size
        scale = min(1.0, max_dim / max(w, h)) if max(w, h) else 1.0
        if scale < 1.0:
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.BILINEAR)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality)
        return buf.getvalue()
    except Exception:
        return data


_VISION_PROMPT = (
    "You are an advanced multimodal document and image analyst. Produce a thorough, "
    "structured analysis of this image that a retrieval system will index and an AI "
    "assistant will later rely on to answer questions about it — it is the ONLY "
    "information the assistant will have, so be exhaustive and precise. Structure your "
    "answer with these exact sections:\n\n"
    "DESCRIPTION: What the image shows in real detail — scene, subject, layout, style, "
    "colors, composition. If it's a UI/screenshot, describe the interface and its state.\n\n"
    "TEXT: Transcribe verbatim any visible text, numbers, labels, code, headings, or UI "
    "copy exactly as shown, preserving structure/line breaks where it matters. Write "
    "'(none)' if there is no legible text.\n\n"
    "DATA: If this is a chart, graph, table, diagram, screenshot of code, or any "
    "structured/data visual, extract the concrete data points, axis labels, values, "
    "code contents, or relationships precisely. Write '(not applicable)' if this isn't "
    "a data visual.\n\n"
    "NOTABLE DETAILS: Anything unusual, important, or likely to be asked about — "
    "errors shown, anomalies, brand/product names, dates, quantities, warnings, etc.\n\n"
    "Do not describe people in a way that could identify a specific real individual by "
    "name unless the name is itself visibly printed in the image (e.g. a name tag or "
    "caption)."
)


async def _analyze_image(data: bytes, filename: str) -> dict:
    """Runs a vision-capable model over the image for a deep, structured analysis.
    Falls back to lightweight Pillow-based metadata if no NVIDIA_API_KEY is
    configured or the vision call fails for any reason, so an upload never
    hard-fails just because live vision analysis wasn't available."""
    width = height = None
    fmt = None
    if Image is not None:
        try:
            img = Image.open(io.BytesIO(data))
            width, height = img.size
            fmt = img.format
        except Exception:
            pass

    if not os.getenv("NVIDIA_API_KEY") or (os.getenv("NVIDIA_API_KEY") or "").strip().lower() in ("demo", ""):
        note = "[Demo mode — no NVIDIA_API_KEY configured, so this image was not analyzed by a vision model.]"
        if width:
            note += f" File appears to be a {fmt or 'image'} image, {width}x{height}px."
        return {"analysis": note, "used_vision_model": False}

    try:
        small = _downscale_image(data)
        b64 = base64.b64encode(small).decode("utf-8")
        llm = ChatNVIDIA(model=VISION_MODEL, temperature=0.2, max_completion_tokens=1500, timeout=90)
        message = HumanMessage(content=[
            {"type": "text", "text": _VISION_PROMPT},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}},
        ])
        result = await llm.ainvoke([message])
        text = _coerce_text(getattr(result, "content", "") or "").strip()
        return {"analysis": text or "[Vision model returned no analysis for this image.]", "used_vision_model": True}
    except Exception as exc:
        print(f"[rag_engine] Vision analysis failed for {filename}: {exc}")
        fallback = f"[Vision analysis unavailable right now ({exc.__class__.__name__}).]"
        if width:
            fallback += f" File is a {fmt or 'image'} image, {width}x{height}px."
        return {"analysis": fallback, "used_vision_model": False}


# Nemotron Parse's required task prompt. Per NVIDIA's own NIM docs this model
# accepts ONLY an image plus this exact kind of control-token string as the
# text turn — no system prompt, no extra instructions, no conversation, no
# tool calls, no follow-up questions (all unsupported by the model itself).
# <predict_bbox><predict_classes><output_markdown> asks for layout-tagged,
# markdown-formatted text; <predict_no_text_in_pic> lets it cleanly say
# "no text here" for a photo instead of inventing something.
_NEMOTRON_PARSE_PROMPT = "<predict_bbox><predict_classes><output_markdown><predict_no_text_in_pic>"

# Nemotron Parse interleaves its actual output with layout/bbox/class control
# tokens, e.g. <class_Title>, <bbox>...</bbox> — none of that is documented
# down to an exact grammar, so rather than parsing it precisely this just
# strips anything that looks like one of the model's own angle-bracket
# tokens and collapses the leftover whitespace, keeping only the readable
# text/markdown content.
_PARSE_TAG_RE = re.compile(r"<[^>\n]{1,40}>")


async def _parse_document_with_nemotron(data: bytes, filename: str) -> Optional[str]:
    """Best-effort OCR/document-structure extraction via nvidia/nemotron-parse-2.0.

    This is purely additive: it runs next to _analyze_image() (never instead
    of it) for every image upload, regardless of which chat model the user
    has selected, and is silently skipped on any failure — a document upload
    should never fail just because this extra pass didn't come back."""
    if not os.getenv("NVIDIA_API_KEY") or (os.getenv("NVIDIA_API_KEY") or "").strip().lower() in ("demo", ""):
        return None
    try:
        small = _downscale_image(data, max_dim=PARSE_IMAGE_MAX_DIM, quality=90)
        b64 = base64.b64encode(small).decode("utf-8")
        llm = ChatNVIDIA(model=PARSE_MODEL, temperature=0.0, max_completion_tokens=3000, timeout=90)
        # Single HumanMessage, no system prompt — Nemotron Parse doesn't
        # support conversational instruction-following, so anything more
        # than image + task prompt here would be ignored at best.
        message = HumanMessage(content=[
            {"type": "text", "text": _NEMOTRON_PARSE_PROMPT},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}},
        ])
        result = await llm.ainvoke([message])
        raw = _coerce_text(getattr(result, "content", "") or "").strip()
        cleaned = _PARSE_TAG_RE.sub(" ", raw)
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
        return cleaned or None
    except Exception as exc:
        print(f"[rag_engine] Nemotron Parse extraction failed for {filename}: {exc}")
        return None


# ---------------------------------------------------------------------------
# Chunking — paragraph-aware, with a small char overlap carried across chunk
# boundaries so retrieval doesn't lose context that straddles a split point.
# ---------------------------------------------------------------------------
def _chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    paras = re.split(r"\n\s*\n", text)
    raw_chunks: List[str] = []
    buf = ""
    for para in paras:
        para = para.strip()
        if not para:
            continue
        if len(para) > size:
            if buf:
                raw_chunks.append(buf)
                buf = ""
            step = max(1, size - overlap)
            for i in range(0, len(para), step):
                raw_chunks.append(para[i:i + size])
            continue
        if len(buf) + len(para) + 2 <= size:
            buf = f"{buf}\n\n{para}" if buf else para
        else:
            if buf:
                raw_chunks.append(buf)
            buf = para
    if buf:
        raw_chunks.append(buf)

    if not raw_chunks:
        return [text[:size]]

    overlapped = [raw_chunks[0]]
    for i in range(1, len(raw_chunks)):
        tail = raw_chunks[i - 1][-overlap:]
        overlapped.append(f"{tail}\n{raw_chunks[i]}")
    return overlapped


# ---------------------------------------------------------------------------
# Embeddings — a real NVIDIA embedding model when configured, otherwise a
# deterministic local hashing vectorizer so retrieval still works with no
# API key (same spirit as the rest of the app's demo-mode fallbacks).
# ---------------------------------------------------------------------------
_embedder_cache: Dict[str, Optional[NVIDIAEmbeddings]] = {}
_WORD_RE = re.compile(r"[A-Za-z0-9_]+")


def _get_embedder() -> Optional[NVIDIAEmbeddings]:
    api_key = (os.getenv("NVIDIA_API_KEY") or "").strip()
    if not api_key or api_key.lower() == "demo":
        return None
    if NVIDIA_EMBED_MODEL in _embedder_cache:
        return _embedder_cache[NVIDIA_EMBED_MODEL]
    try:
        embedder = NVIDIAEmbeddings(model=NVIDIA_EMBED_MODEL)
        _embedder_cache[NVIDIA_EMBED_MODEL] = embedder
        return embedder
    except Exception as exc:
        print(f"[rag_engine] Could not initialize NVIDIAEmbeddings({NVIDIA_EMBED_MODEL}): {exc}")
        _embedder_cache[NVIDIA_EMBED_MODEL] = None
        return None


def _local_embed(text: str, dim: int = EMBED_DIM_FALLBACK) -> List[float]:
    """Deterministic hashed bag-of-words/bigrams vector. No network calls, no
    external state — always available, used whenever a real embedding model
    isn't configured or a live embedding call fails."""
    vec = np.zeros(dim, dtype=np.float32)
    words = _WORD_RE.findall((text or "").lower())
    if not words:
        return vec.tolist()
    for w in words:
        idx = int(hashlib.md5(w.encode("utf-8")).hexdigest(), 16) % dim
        vec[idx] += 1.0
    for a, b in zip(words, words[1:]):
        idx = int(hashlib.md5(f"{a}_{b}".encode("utf-8")).hexdigest(), 16) % dim
        vec[idx] += 0.5
    norm = float(np.linalg.norm(vec)) or 1.0
    return (vec / norm).tolist()


async def _embed_texts(texts: List[str]) -> Tuple[List[List[float]], str]:
    embedder = _get_embedder()
    if embedder is not None:
        try:
            vectors = await asyncio.to_thread(embedder.embed_documents, texts)
            return vectors, f"nvidia:{NVIDIA_EMBED_MODEL}"
        except Exception as exc:
            print(f"[rag_engine] NVIDIA embedding call failed, falling back to local embeddings: {exc}")
    return [_local_embed(t) for t in texts], "local"


async def _embed_query(query: str, space: str) -> List[float]:
    if space.startswith("nvidia:"):
        embedder = _get_embedder()
        if embedder is not None:
            try:
                return await asyncio.to_thread(embedder.embed_query, query)
            except Exception as exc:
                print(f"[rag_engine] NVIDIA query embedding failed, falling back to local: {exc}")
    return _local_embed(query)


def _lexical_overlap(query: str, text: str) -> float:
    """Cheap Jaccard-style overlap on word sets — used as a small hybrid boost
    so exact terms (names, error strings, filenames, numbers) that a semantic
    embedding space might under-weight still surface reliably."""
    qw = set(_WORD_RE.findall((query or "").lower()))
    tw = set(_WORD_RE.findall((text or "").lower()))
    if not qw or not tw:
        return 0.0
    inter = len(qw & tw)
    if inter == 0:
        return 0.0
    return inter / max(1, len(qw))


# ---------------------------------------------------------------------------
# Per-session store. Operates directly on the `session: dict` that app.py
# already owns and passes around — no separate global registry to keep in
# sync, and files/chunks are freed automatically when a session is cleared.
# ---------------------------------------------------------------------------
def _files_store(session: dict) -> Dict[str, dict]:
    return session.setdefault("rag_files", {})


def _chunks_store(session: dict) -> List[dict]:
    return session.setdefault("rag_chunks", [])


def has_files(session: dict) -> bool:
    return bool(session.get("rag_files"))


def list_files(session: dict) -> List[dict]:
    files = session.get("rag_files") or {}
    return sorted(
        ({k: v for k, v in rec.items() if not k.startswith("_")} for rec in files.values()),
        key=lambda f: f.get("uploaded_at", ""),
    )


def get_file_content(session: dict, file_id: str) -> Optional[Tuple[bytes, str, str]]:
    """Return raw attachment bytes plus metadata for a session-scoped preview."""
    record = (session.get("rag_files") or {}).get(file_id)
    if not record or record.get("status") == "error":
        return None
    data = record.get("_content")
    if not isinstance(data, (bytes, bytearray)):
        return None
    return bytes(data), record.get("content_type") or "application/octet-stream", record.get("filename") or "download"


async def wait_for_processing(session: dict, attachment_ids: Optional[List[str]] = None, timeout: float = 120.0) -> None:
    """Wait for background indexing only when a response actually needs RAG."""
    deadline = asyncio.get_running_loop().time() + timeout
    allowed = set(attachment_ids or [])
    while asyncio.get_running_loop().time() < deadline:
        files = session.get("rag_files") or {}
        records = [rec for rec in files.values() if not allowed or rec.get("id") in allowed]
        if not any(rec.get("status") == "processing" for rec in records):
            return
        await asyncio.sleep(0.1)


def remove_file(session: dict, file_id: str) -> bool:
    files = session.get("rag_files") or {}
    if file_id not in files:
        return False
    files.pop(file_id, None)
    session["rag_chunks"] = [c for c in (session.get("rag_chunks") or []) if c["file_id"] != file_id]
    session_store.save_rag_state(session)
    return True


def create_upload_record(filename: str, data: bytes, content_type: str, session: dict) -> dict:
    """Create the lightweight record returned immediately by the upload API."""
    files = _files_store(session)
    file_id = uuid.uuid4().hex[:12]
    kind = classify(filename, content_type)
    record: dict = {
        "id": file_id,
        "filename": filename,
        "kind": kind,
        "size": len(data),
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "status": "processing",
        "summary": "",
        "chunk_count": 0,
        "error": None,
        "thumbnail": None,
        "used_vision_model": False,
        "used_document_parser": False,
        "content_type": content_type or "application/octet-stream",
        "_content": data,
    }
    files[file_id] = record
    if len(files) > MAX_FILES_PER_SESSION:
        record.update(status="error", error=f"This session already has {MAX_FILES_PER_SESSION} files attached — remove one before adding more.")
    elif len(data) > MAX_FILE_BYTES:
        record.update(status="error", error=f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)}MB limit.")
    elif kind == "unsupported":
        record.update(status="error", error="Unsupported file type.")
    elif kind == "image":
        try:
            thumb = _downscale_image(data, max_dim=THUMB_MAX_DIM, quality=70)
            record["thumbnail"] = f"data:image/jpeg;base64,{base64.b64encode(thumb).decode('utf-8')}"
        except Exception:
            pass
    return record


async def process_upload(filename: str, data: bytes, content_type: str, session: dict, record: Optional[dict] = None) -> dict:
    """Main entry point: extract -> chunk -> embed -> store one uploaded file.
    Always returns a record dict (never raises) so the API layer can surface
    a clean per-file status even when this particular file failed."""
    files = _files_store(session)
    chunks = _chunks_store(session)

    if record is None:
        record = create_upload_record(filename, data, content_type, session)
    file_id = record["id"]
    kind = record["kind"]
    size = record["size"]

    try:
        if len(files) > MAX_FILES_PER_SESSION:
            record.update(status="error", error=f"This session already has {MAX_FILES_PER_SESSION} files attached — remove one before adding more.")
            return record
        if size > MAX_FILE_BYTES:
            record.update(status="error", error=f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)}MB limit.")
            return record
        if kind == "unsupported":
            record.update(status="error", error="Unsupported file type.")
            return record

        raw_pieces: List[Tuple[Optional[int], str]] = []  # (page_number, text)

        if kind == "image":
            # General scene/description analysis and the Nemotron Parse OCR/
            # document-structure pass run concurrently — two independent
            # NVIDIA vision calls over the same image, combined below. Either
            # one can fail without affecting the other (see their own
            # try/except blocks), and neither is ever surfaced as a "model"
            # the user picks — this is pure backend preprocessing.
            analysis, parsed_text = await asyncio.gather(
                _analyze_image(data, filename),
                _parse_document_with_nemotron(data, filename),
            )
            record["used_vision_model"] = analysis["used_vision_model"]
            record["used_document_parser"] = bool(parsed_text)
            content = analysis["analysis"]
            if parsed_text:
                content += (
                    "\n\n---\nDOCUMENT TEXT EXTRACTION (verbatim OCR via NVIDIA Nemotron "
                    "Parse — trust this over the description above for exact wording, "
                    "numbers, or table contents):\n" + parsed_text
                )
            raw_pieces = [(None, content)]
            record["summary"] = _short_summary(analysis["analysis"])

        elif kind == "pdf":
            pages = await _extract_pdf(data)
            raw_pieces = pages
            preview = " ".join(t for _, t in pages[:2] if t)
            record["summary"] = _short_summary(preview) or f"{len(pages)}-page PDF"

        elif kind == "docx":
            text = await _extract_docx(data)
            raw_pieces = [(None, text)]
            record["summary"] = _short_summary(text)

        else:  # text or code
            text = _extract_plain(data)
            raw_pieces = [(None, text)]
            record["summary"] = _short_summary(text)

        all_pieces: List[Tuple[Optional[int], str]] = []
        for page, text in raw_pieces:
            for piece in _chunk_text(text):
                all_pieces.append((page, piece))

        if not all_pieces:
            record.update(status="error", error="No extractable content found in this file.")
            return record

        # --- REAL FIX (see conversation) --------------------------------
        # Publish the extracted chunks — vector=None, not yet embedded —
        # and flip status to "ready" right here, BEFORE calling the
        # embedding model. Content (the vision analysis / extracted text)
        # is already fully available at this point; embedding is only
        # needed for *semantic search across many older files later*, not
        # for answering about the file you just attached in this same
        # turn (build_context() takes a direct, non-semantic path for
        # freshly-attached files — see get_all_text_for_files()). Splitting
        # these means a chat reply about a just-uploaded file only ever
        # waits on extraction (one vision call for images, near-instant
        # local parsing for text/pdf/docx) instead of extraction PLUS an
        # embedding network round-trip, which is what made this app feel
        # slower than ChatGPT for "here's a file, tell me about it".
        new_chunks = [
            {
                "id": uuid.uuid4().hex[:12],
                "file_id": file_id,
                "filename": filename,
                "kind": kind,
                "page": page,
                "text": text,
                "vector": None,
                "space": "pending",
            }
            for page, text in all_pieces
        ]
        chunks.extend(new_chunks)
        record.update(status="ready", chunk_count=len(all_pieces))

        # Embedding continues in the background (this coroutine keeps
        # running as the tracked upload task) purely to make this file
        # semantically searchable on LATER turns. Mutating the same dicts
        # already in `chunks` in place means any concurrent reader sees
        # the upgrade the instant it lands, with no extra bookkeeping.
        texts = [c["text"] for c in new_chunks]
        vectors, space = await _embed_texts(texts)
        for chunk, vector in zip(new_chunks, vectors):
            chunk["vector"] = vector
            chunk["space"] = space

        return record

    except Exception as exc:
        traceback.print_exc()
        record.update(status="error", error=str(exc)[:300])
        return record
    finally:
        # Mirror rag_files/rag_chunks to Redis (no-op if REDIS_URL isn't set) so
        # retrieval survives a process restart/spin-down instead of silently
        # going blank — see session_store.py.
        session_store.save_rag_state(session)


# ---------------------------------------------------------------------------
# Retrieval
# ---------------------------------------------------------------------------
async def retrieve(session: dict, query: str, top_k: int = TOP_K_DEFAULT,
                    attachment_ids: Optional[List[str]] = None) -> List[dict]:
    """Hybrid semantic + lexical retrieval over this session's stored chunks.
    Vectorized with numpy per embedding-space group for speed even with a
    large number of stored chunks."""
    all_chunks = session.get("rag_chunks") or []
    if attachment_ids:
        allowed = set(attachment_ids)
        all_chunks = [c for c in all_chunks if c["file_id"] in allowed]
    if not all_chunks or not (query or "").strip():
        return []

    by_space: Dict[str, List[dict]] = {}
    for c in all_chunks:
        by_space.setdefault(c.get("space") or "local", []).append(c)

    scored: List[Tuple[float, dict]] = []
    # Chunks whose background embedding hasn't landed yet ("pending", or any
    # vector-less entry) can't get a semantic score — fall back to lexical
    # overlap alone for those instead of dropping them from retrieval.
    pending_group = by_space.pop("pending", [])
    pending_group += [c for group in by_space.values() for c in group if c.get("vector") is None]
    if pending_group:
        lex = np.array([_lexical_overlap(query, c["text"]) for c in pending_group], dtype=np.float32)
        scored.extend(zip(lex.tolist(), pending_group))

    for space, group in by_space.items():
        group = [c for c in group if c.get("vector") is not None]
        if not group:
            continue
        qvec = np.array(await _embed_query(query, space), dtype=np.float32)
        mat = np.array([c["vector"] for c in group], dtype=np.float32)
        q_norm = float(np.linalg.norm(qvec)) or 1e-9
        row_norms = np.linalg.norm(mat, axis=1)
        row_norms[row_norms == 0] = 1e-9
        sims = (mat @ qvec) / (row_norms * q_norm)
        lex = np.array([_lexical_overlap(query, c["text"]) for c in group], dtype=np.float32)
        hybrid = sims * 0.82 + lex * 0.18
        scored.extend(zip(hybrid.tolist(), group))

    scored.sort(key=lambda pair: pair[0], reverse=True)
    top = scored[:max(1, top_k)]
    return [{"score": round(float(score), 4), **{k: v for k, v in c.items() if k != "vector"}} for score, c in top]


def format_context(chunks: List[dict], max_chars_per_chunk: int = MAX_CONTEXT_CHARS_PER_CHUNK) -> str:
    if not chunks:
        return ""
    blocks = []
    for c in chunks:
        loc = f", page {c['page']}" if c.get("page") else ""
        text = c["text"]
        if len(text) > max_chars_per_chunk:
            text = text[:max_chars_per_chunk].rstrip() + "…"
        kind_tag = " (image analysis)" if c.get("kind") == "image" else ""
        blocks.append(f"[FILE: {c['filename']}{loc}{kind_tag} — relevance {c['score']:.2f}]\n{text}")
    return "\n\n---\n\n".join(blocks)


def get_all_text_for_files(session: dict, file_ids: List[str], max_chars_per_file: int = 6000) -> str:
    """Direct, non-semantic context for files attached to THIS turn — no
    embed-and-search round trip. If the user just uploaded a file and is
    asking about it in the same message, it's already known to be relevant;
    semantic retrieval exists to find the right needle in a haystack of
    *older* uploads, not to re-confirm the one file just handed to you."""
    if not file_ids:
        return ""
    wanted = set(file_ids)
    chunks = [c for c in (session.get("rag_chunks") or []) if c["file_id"] in wanted]
    if not chunks:
        return ""
    by_file: Dict[str, List[dict]] = {}
    for c in chunks:
        by_file.setdefault(c["file_id"], []).append(c)
    blocks = []
    for cs in by_file.values():
        cs.sort(key=lambda c: c.get("page") or 0)
        text = "\n\n".join(c["text"] for c in cs)
        if len(text) > max_chars_per_file:
            text = text[:max_chars_per_file].rstrip() + "…"
        kind_tag = " (image analysis)" if cs[0].get("kind") == "image" else ""
        blocks.append(f"[FILE: {cs[0]['filename']}{kind_tag}]\n{text}")
    return "\n\n---\n\n".join(blocks)


async def build_context(session: dict, query: str, attachment_ids: Optional[List[str]] = None,
                         top_k: int = TOP_K_DEFAULT) -> Tuple[str, List[dict]]:
    """Convenience wrapper: retrieve + format in one call. Returns
    (context_text, chunks_used) — chunks_used is [] and context_text is ""
    whenever the session has no uploaded files, so callers can skip the RAG
    prompt block entirely without an extra has_files() check.

    Files named in attachment_ids (just attached this turn) go straight in
    via get_all_text_for_files() — no embedding/search needed. Any other
    files already in the session (older uploads not referenced this turn)
    still go through normal semantic retrieve() so cross-turn recall keeps
    working."""
    if not session.get("rag_chunks"):
        return "", []
    if attachment_ids:
        direct_text = get_all_text_for_files(session, attachment_ids)
        other_ids = {c["file_id"] for c in session["rag_chunks"]} - set(attachment_ids)
        chunks, retrieved_text = [], ""
        if other_ids:
            chunks = await retrieve(session, query, top_k=top_k, attachment_ids=list(other_ids))
            retrieved_text = format_context(chunks)
        combined = "\n\n---\n\n".join(x for x in (direct_text, retrieved_text) if x)
        return combined, chunks
    chunks = await retrieve(session, query, top_k=top_k)
    return format_context(chunks), chunks
